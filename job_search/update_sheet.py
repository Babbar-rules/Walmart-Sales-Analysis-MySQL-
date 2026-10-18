"""Step 3: append Claude's selected NEW jobs to Harpreet_Job_Matches.xlsx after verifying every link.

Input: data/selected.json - list of objects with keys
  id (from shortlist.json), company, location, mode, experience, score (0-100), description, why
Also:
  * re-checks every job already in the sheet and marks closed ones (never duplicates them),
  * records every reviewed shortlist URL in seen_jobs.json so it is never re-added or re-reviewed,
  * refreshes the "Companies Searched" sheet.
"""
import datetime as dt, json, os, re, sys, urllib.error, urllib.request
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from fetch_jobs import DATA, HERE, get

XLSX = os.path.join(os.path.dirname(HERE), "Harpreet_Job_Matches.xlsx")
TODAY = dt.date.today().isoformat()
BROWSER = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126 Safari/537.36"}
HDR = PatternFill("solid", fgColor="1F4E78"); THIN = Side(style="thin", color="BFBFBF")


def http_ok(u):
    try: return urllib.request.urlopen(urllib.request.Request(u, headers=BROWSER), timeout=30).status == 200
    except Exception: return False


def live(u):
    """True/False when the careers API can confirm the posting is open; None when there is no API to ask."""
    if "ashbyhq.com" in u:
        s, jid = u.split("/")[3:5]
        d = get(f"https://api.ashbyhq.com/posting-api/job-board/{s}")
        return None if d is None else any(jid in x["jobUrl"] for x in d.get("jobs", []))
    if "jobs.lever.co" in u:
        s, jid = u.split("/")[3:5]
        return bool((get(f"https://api.lever.co/v0/postings/{s}/{jid}", tries=1) or {}).get("id"))
    if "myworkdayjobs.com" in u:
        m = re.match(r"(https://([^.]+)\.wd\d+\.myworkdayjobs\.com)/([^/]+)(/job/.*)", u)
        d = m and get(f"{m[1]}/wday/cxs/{m[2]}/{m[3]}{m[4]}", tries=1)
        return bool(d and d.get("jobPostingInfo", {}).get("canApply", True))
    if "amazon.jobs" in u:
        try: body = urllib.request.urlopen(urllib.request.Request(u, headers=BROWSER), timeout=30).read().decode("utf8", "ignore")
        except Exception: return False
        return "no longer available" not in body.lower()
    return None


def verdict(u):
    ok, lv = http_ok(u), live(u)
    if ok and lv is not False: return True, "✅ Working – HTTP 200, posting live"
    return False, "❌ Closed / link not working"


def role_key(company, title):
    return (re.sub(r"\W+", "", str(company)).lower(), re.sub(r"\W+", "", str(title)).lower())


def style_row(ws, r):
    for c in ws[r]:
        c.alignment = Alignment(wrap_text=True, vertical="top"); c.border = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def main():
    seen = json.load(open(os.path.join(HERE, "seen_jobs.json")))
    shortlist = {j["id"]: j for j in json.load(open(os.path.join(DATA, "shortlist.json")))}
    sel_path = os.path.join(DATA, "selected.json")
    selected = json.load(open(sel_path)) if os.path.exists(sel_path) else []
    wb = load_workbook(XLSX); ws = wb["Matching Jobs"]
    if ws.cell(1, 12).value != "Date Added":
        h = ws.cell(1, 12, "Date Added"); h.font = Font(bold=True, color="FFFFFF"); h.fill = HDR
        for r in range(2, ws.max_row + 1): ws.cell(r, 12, "2026-09-29")
        ws.column_dimensions["L"].width = 13
    ws.cell(1, 11).value = "Link Check (last run)"

    # 1. Re-check jobs already in the sheet; mark closed ones instead of removing/duplicating them.
    in_sheet = set(); roles = set(); closed = 0
    for r in range(2, ws.max_row + 1):
        u = ws.cell(r, 10).value
        if not u: continue
        in_sheet.add(u); roles.add(role_key(ws.cell(r, 2).value, ws.cell(r, 3).value))
        ok, msg = verdict(u)
        ws.cell(r, 11).value = f"{msg} ({TODAY})"
        if not ok:
            closed += 1
            for c in ws[r]: c.font = Font(color="999999", strike=True) if c.column != 11 else Font(color="C00000")

    # 2. Append only verified NEW selections.
    added, rejected_links = [], []
    for s in sorted(selected, key=lambda s: -s["score"]):
        j = shortlist[s["id"]]
        if j["url"] in in_sheet or j["url"] in seen: continue
        if role_key(s["company"], j["title"]) in roles: continue  # same role re-posted under a new link
        ok, msg = verdict(j["url"])
        if not ok: rejected_links.append(j["url"]); continue
        n = ws.max_row
        ws.append([n, s["company"], j["title"].strip(), s["location"], s["mode"], s["experience"], s["score"] / 100,
                   s["description"], s["why"], j["url"], f"{msg} ({TODAY})", TODAY])
        r = ws.max_row; style_row(ws, r)
        ws.cell(r, 10).hyperlink = j["url"]; ws.cell(r, 10).font = Font(color="0563C1", underline="single")
        ws.cell(r, 7).number_format = "0%"
        ws.cell(r, 7).fill = PatternFill("solid", fgColor="E2EFDA" if s["score"] >= 85 else "FFF2CC" if s["score"] >= 75 else "FCE4D6")
        ws.cell(r, 12).fill = PatternFill("solid", fgColor="DDEBF7")
        in_sheet.add(j["url"]); roles.add(role_key(s["company"], j["title"])); added.append(f"{s['score']}% {s['company']} – {j['title'].strip()}")
    ws.auto_filter.ref = ws.dimensions

    # 3. Remember everything reviewed this run (selected or not) so it is never re-added or re-reviewed.
    seen = sorted(set(seen) | in_sheet | {j["url"] for j in shortlist.values()})
    json.dump(seen, open(os.path.join(HERE, "seen_jobs.json"), "w"), indent=0)

    # 4. Refresh the Companies Searched sheet.
    comps = json.load(open(os.path.join(DATA, "companies.json")))
    names = {}
    for r in range(2, ws.max_row + 1):
        u = ws.cell(r, 10).value or ""
        names[ws.cell(r, 2).value] = names.get(ws.cell(r, 2).value, 0) + 1
    if "Companies Searched" in wb.sheetnames: del wb["Companies Searched"]
    ws2 = wb.create_sheet("Companies Searched", 1)
    ws2.append(["Company (board)", "Official Careers Page", "Source", "Openings Scanned (last run)", "Jobs in Sheet", "Last Run"])
    for c in ws2[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR; c.alignment = Alignment(wrap_text=True, horizontal="center")
    for c in sorted(comps, key=lambda c: c["slug"].lower()):
        cnt = sum(v for k, v in names.items() if k and (c["slug"].lower() in k.lower().replace(" ", "").replace(".", "") or k.lower().replace(" ", "") in c["slug"].lower()))
        ws2.append([c["slug"], c["board"], c["source"], c["openings"], cnt, TODAY])
        ws2.cell(ws2.max_row, 2).hyperlink = c["board"]; ws2.cell(ws2.max_row, 2).font = Font(color="0563C1", underline="single")
    for col, w in zip("ABCDEF", [24, 62, 16, 20, 14, 12]): ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    wb.save(XLSX)

    summary = dict(date=TODAY, added=added, closed_existing=closed, dropped_bad_links=rejected_links, total_in_sheet=len(in_sheet))
    json.dump(summary, open(os.path.join(DATA, "last_run.json"), "w"), indent=1)
    print(json.dumps(summary, indent=1, ensure_ascii=False))


if __name__ == "__main__":
    sys.exit(main())
